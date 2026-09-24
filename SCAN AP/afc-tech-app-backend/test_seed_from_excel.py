#!/usr/bin/env python3
"""Create a tiny survey workbook, seed twice, and confirm AHUs update instead of duplicating."""
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from flask import Flask
from openpyxl import Workbook

from db import db
import models  # noqa: F401 — register tables
from models import Hospital, AHU, Filter
from datetime import date

from seed_from_excel import (
    normalize_filter_size,
    part_match_keys,
    seed_from_excel,
    serialize_seed_stats,
)


HEADERS = [
    "AHU NO.",
    "LOCATION",
    "STAGE",
    "FILTER SIZE",
    "FREQUENCY",
    "QUANTITY",
    "BUILDING",
    "FLOOR/AREA",
    "PART NUMBER",
    "FILTER TYPE",
    "DATE OF REPLACEMENT",
]


def write_workbook(path, hospital, rows, sheet="MAIN BUILDING"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["B2"] = hospital
    for col, header in enumerate(HEADERS, 1):
        ws.cell(5, col, header)
    for r_i, row in enumerate(rows, start=6):
        for c_i, val in enumerate(row, 1):
            ws.cell(r_i, c_i, val)
    wb.save(path)


def sample_rows(qty_pre=2, location="Roof"):
    return [
        ["AH-1", location, "Pre", "12x12x1", "90 days", qty_pre, "Main", "2nd", "PN-1", "Pleat", None],
        ["AH-1", location, "Final", "24x24x2", "180 days", 1, "Main", "2nd", "PN-2", "Bag", None],
    ]


def make_app(db_path):
    app = Flask(__name__)
    app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + db_path.replace("\\", "/")
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    db.init_app(app)
    with app.app_context():
        db.create_all()
    return app


def assert_eq(actual, expected, msg):
    if actual != expected:
        raise AssertionError(f"{msg}: expected {expected!r}, got {actual!r}")


def main():
    payload = serialize_seed_stats(
        {
            "hospital": "X",
            "hospital_id": 1,
            "ahus": {3, 1},
            "ahus_created": 0,
            "ahus_updated": 2,
            "filters_upserted": 4,
        },
        dry_run=True,
    )
    assert_eq(payload["dry_run"], True, "dry_run flag")
    assert_eq(payload["ahus_touched"], 2, "set length is JSON-safe")

    db_fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(db_fd)
    app = make_app(db_path)
    with app.app_context():
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            write_workbook(path, "Foothill", sample_rows())
            first = seed_from_excel(path)
            assert_eq(first["hospital_created"], True, "first pass creates hospital")
            assert_eq(first["ahus_created"], 1, "first pass creates AHU")
            assert_eq(first["ahus_updated"], 0, "first pass no updates")
            assert_eq(first["filters_upserted"], 2, "two filters")
            assert_eq(Hospital.query.count(), 1, "one hospital")
            assert_eq(AHU.query.count(), 1, "one AHU")
            assert_eq(Filter.query.count(), 2, "two filters")
            hid = first["hospital_id"]
            ahu_id = AHU.query.first().id
            assert_eq(Filter.query.filter_by(part_number="PN-1").first().quantity, 2, "initial qty")

            write_workbook(
                path,
                "foothill",
                [
                    ["ah-1", "Roof penthouse", "Pre", "12x12x1", "90 days", 5, "MAIN", "2nd", "PN-1", "Pleat", None],
                    ["ah-1", "Roof penthouse", "Final", "24x24x2", "180 days", 1, "MAIN", "2nd", "PN-2", "Bag", None],
                ],
            )
            dry = seed_from_excel(path, dry_run=True)
            assert_eq(dry["dry_run"], True, "dry run flag")
            assert_eq(dry["ahus_created"], 0, "dry run matches existing AHU")
            assert_eq(dry["ahus_updated"], 1, "dry run update")
            assert_eq(AHU.query.count(), 1, "dry run does not commit a second AHU")
            assert_eq(
                Filter.query.filter_by(part_number="PN-1").first().quantity,
                2,
                "dry run does not change qty",
            )

            second = seed_from_excel(path)
            assert_eq(second["hospital_created"], False, "case-insensitive hospital match")
            assert_eq(second["ahus_created"], 0, "update not create")
            assert_eq(second["ahus_updated"], 1, "matched AHU")
            assert_eq(Hospital.query.count(), 1, "still one hospital")
            assert_eq(AHU.query.count(), 1, "still one AHU")
            assert_eq(Filter.query.count(), 2, "still two filters")
            assert_eq(Filter.query.filter_by(part_number="PN-1").first().quantity, 5, "qty updated")
            assert_eq(AHU.query.first().id, ahu_id, "same AHU id")
            assert_eq(AHU.query.first().location, "Roof penthouse", "location updated")

            write_workbook(path, "Wrong Name", sample_rows(qty_pre=8, location="Penthouse"))
            third = seed_from_excel(path, hospital_id=hid)
            assert_eq(third["hospital_created"], False, "hospital_id override does not create")
            assert_eq(Hospital.query.count(), 1, "B2 mismatch did not spawn a hospital")
            assert_eq(AHU.query.count(), 1, "still one AHU after override")
            assert_eq(Filter.query.filter_by(part_number="PN-1").first().quantity, 8, "qty updated via override")
            assert any("B2" in w for w in third["warnings"]), "warn when B2 disagrees with selected hospital"

            assert_eq(normalize_filter_size("24x24x2 HV"), "24x24x2", "strip HV from size")
            assert_eq(normalize_filter_size("24x24x12 FF"), "24x24x12", "strip FF from size")
            assert "F8V42412GWBB" in part_match_keys("F8V424-GWBB", "24x24x12"), "part + depth"
            assert part_match_keys("F8V424-GWBB", "24x24x12") & part_match_keys(
                "F8V42412-GWBB", "24x24x12"
            ), "catalog part matches stored part with depth"

            write_workbook(
                path,
                "Foothill",
                [
                    ["AH-1", "Penthouse", "PRE", "24x24x2 HV", "90 days", 32, "MAIN", "Roof", "HVP24242", "Pleated", date(2026, 8, 25)],
                    ["AH-1", "Penthouse", "FINAL", "24x24x12 FF", "365 days", 16, "MAIN", "Roof", "F8V424-GWBB", "V-Bank", date(2026, 5, 5)],
                ],
            )
            # Pretend the app already had these rows under slightly different labels
            ahu = AHU.query.first()
            Filter.query.delete()
            db.session.add(Filter(
                ahu_id=ahu.id, phase="PRE", part_number="HVP24242", size="24x24x2",
                quantity=32, frequency_days=90, last_service_date=date(2025, 10, 22), is_active=True,
            ))
            db.session.add(Filter(
                ahu_id=ahu.id, phase="PRE", part_number="HVP24242", size="24x24x2 HV",
                quantity=32, frequency_days=90, last_service_date=date(2026, 8, 25), is_active=True,
            ))
            db.session.add(Filter(
                ahu_id=ahu.id, phase="FINAL", part_number="F8V42412-GWBB", size="24x24x12",
                quantity=16, frequency_days=365, last_service_date=date(2026, 5, 5), is_active=True,
            ))
            db.session.commit()
            assert_eq(Filter.query.filter_by(is_active=True).count(), 3, "setup: 3 active filters")

            fourth = seed_from_excel(path, hospital_id=hid)
            active = Filter.query.filter_by(ahu_id=ahu.id, is_active=True).all()
            assert_eq(len(active), 2, "survey suffixes do not create extra filters")
            pre = next(f for f in active if normalize_filter_size(f.size) == "24x24x2")
            final = next(f for f in active if normalize_filter_size(f.size) == "24x24x12")
            assert_eq(pre.last_service_date, date(2026, 8, 25), "PRE date updated from survey")
            assert_eq(pre.quantity, 32, "PRE qty kept")
            assert_eq(final.part_number, "F8V42412-GWBB", "keep original part number")
            assert_eq(final.last_service_date, date(2026, 5, 5), "FINAL date from survey")
            assert_eq(fourth["ahus_created"], 0, "AHU still matched")
            assert_eq(Filter.query.filter_by(ahu_id=ahu.id).count(), 2, "unused extra PRE row deleted")
        finally:
            try:
                os.unlink(path)
            except OSError:
                pass
            try:
                os.unlink(db_path)
            except OSError:
                pass

    print("seed_from_excel tests passed")


if __name__ == "__main__":
    main()
